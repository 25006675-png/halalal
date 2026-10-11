"""
HalalReady backend: endpoints the dashboard calls (CONTRACT.md §7).
Owner actions (approve, reject, resolve, upload, date jump) live here too.
"""

import asyncio
import json
import shutil
import uuid
from datetime import datetime, timedelta
from typing import Literal, Optional

from fastapi import APIRouter, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field
from sse_starlette.sse import EventSourceResponse

from . import calculations as calc
from . import messaging, pack, state
from .database import (
    AGENTS, MYT, UPLOAD_DIR, ZONES, connect, fetch_all, fetch_one, fingerprint, insert, now_iso, reset_db, update,
)
from .seed import seed

router = APIRouter(tags=["dashboard"])


async def _business(db, business_id: int) -> dict:
    biz = await fetch_one(db, "SELECT * FROM businesses WHERE id = ?", (business_id,))
    if not biz:
        raise HTTPException(404, f"Business {business_id} not found")
    return biz


# ── health, clock, reset ─────────────────────────────────────────────

@router.get("/health")
async def health():
    return {"ok": True}


@router.get("/clock")
async def get_clock():
    async with connect() as db:
        override = await state.get_today_override(db)
        today = await state.get_today(db)
    return {"today": today.isoformat(), "override": override,
            "real_today": calc.effective_today(None).isoformat()}


class SetClock(BaseModel):
    today: Optional[str] = Field(default=None, pattern=r"^\d{4}-\d{2}-\d{2}$")
    add_days: Optional[int] = None
    business_id: int = 1


@router.put("/clock")
async def set_clock(req: SetClock):
    """Jump the date for the Upkeep demo. Wakes Renewal Watch for anything now due."""
    async with connect() as db:
        if req.today:
            new_today = req.today
        elif req.add_days is not None:
            new_today = ((await state.get_today(db)) + timedelta(days=req.add_days)).isoformat()
        else:
            raise HTTPException(422, "Give today or add_days")
        await state.set_today_override(db, new_today)
        await db.commit()
        await state.post_event(db, req.business_id, "system", f"Date moved to {new_today}", "info")
        triggered = await state.upkeep_check(db, req.business_id)
        biz = await _business(db, req.business_id)
    alert_channel = None
    if triggered:
        names = ", ".join(t["title"] for t in triggered)
        alert_channel = await messaging.notify_owner(
            biz, f"HalalReady: renewal coming up. {names}. Renewal Watch has started; "
                 f"you'll see messages to approve in the app.")
    return {"today": new_today, "override": new_today,
            "triggered_tasks": [state.task_out(t) for t in triggered], "owner_alert": alert_channel}


@router.delete("/clock")
async def reset_clock():
    async with connect() as db:
        await state.set_today_override(db, None)
    return await get_clock()


@router.post("/admin/reset")
async def admin_reset():
    """One-command reset for the demo: empty database, uploads cleared, sample business seeded."""
    await reset_db()
    if UPLOAD_DIR.exists():
        shutil.rmtree(UPLOAD_DIR, ignore_errors=True)
    await seed()
    return {"reset": True}


# ── runs ─────────────────────────────────────────────────────────────

class StartRun(BaseModel):
    mode: Literal["prepare", "upkeep"] = "prepare"
    instruction: Optional[str] = None   # the owner's one prompt; becomes the Coordinator's first task


@router.post("/business/{business_id}/runs")
async def start_run(business_id: int, req: StartRun = StartRun()):
    now = now_iso()
    async with connect() as db:
        await _business(db, business_id)
        await db.execute("UPDATE runs SET finished_at = ? WHERE business_id = ? AND finished_at IS NULL",
                         (now, business_id))
        run_id = await insert(db, "runs", {"business_id": business_id, "mode": req.mode, "started_at": now})
        await db.commit()
        if req.instruction:
            await state.create_task(db, business_id, "owner", "coordinator", req.instruction.strip()[:200],
                                    run_id=run_id, event_label=req.instruction.strip()[:120])
    return {"id": run_id, "business_id": business_id, "mode": req.mode, "started_at": now}


@router.post("/runs/{run_id}/finish")
async def finish_run(run_id: int):
    async with connect() as db:
        await db.execute("UPDATE runs SET finished_at = ? WHERE id = ? AND finished_at IS NULL", (now_iso(), run_id))
        return await fetch_one(db, "SELECT * FROM runs WHERE id = ?", (run_id,))


@router.get("/runs/{run_id}")
async def get_run(run_id: int):
    async with connect() as db:
        run = await fetch_one(db, "SELECT * FROM runs WHERE id = ?", (run_id,))
        if not run:
            raise HTTPException(404, "Run not found")
        events = await fetch_all(db, "SELECT * FROM events WHERE run_id = ? ORDER BY id", (run_id,))
        tasks = await fetch_all(db, "SELECT * FROM tasks WHERE run_id = ? ORDER BY id", (run_id,))
    return {"run": run, "events": events, "tasks": [state.task_out(t) for t in tasks]}


# ── events and the live stream ───────────────────────────────────────

@router.get("/business/{business_id}/events")
async def list_events(business_id: int, after_id: int = 0, run_id: Optional[int] = None,
                      agent: Optional[str] = None, limit: int = 200):
    sql = "SELECT * FROM events WHERE business_id = ? AND id > ?"
    params: list = [business_id, after_id]
    if run_id is not None:
        sql += " AND run_id = ?"
        params.append(run_id)
    if agent:
        sql += " AND agent = ?"
        params.append(agent)
    async with connect() as db:
        return await fetch_all(db, sql + " ORDER BY id LIMIT ?", (*params, min(limit, 1000)))


async def _stream(request: Request, business_id: int, run_id: Optional[int], after_id: Optional[int]):
    """Backlog first, then live events. Each message: id = event id, data = event JSON."""
    last_header = request.headers.get("last-event-id")
    if after_id is None and last_header and last_header.isdigit():
        after_id = int(last_header)
    queue = state.hub.subscribe()

    def wanted(ev: dict) -> bool:
        return ev["business_id"] == business_id and (run_id is None or ev["run_id"] == run_id)

    async def generator():
        last_sent = after_id or 0
        try:
            async with connect() as db:
                if after_id is None and run_id is None:
                    # No position given: send the latest 100 events for context.
                    backlog = await fetch_all(
                        db, "SELECT * FROM (SELECT * FROM events WHERE business_id = ? ORDER BY id DESC LIMIT 100) "
                            "ORDER BY id", (business_id,))
                else:
                    sql = "SELECT * FROM events WHERE business_id = ? AND id > ?"
                    params = [business_id, last_sent]
                    if run_id is not None:
                        sql += " AND run_id = ?"
                        params.append(run_id)
                    backlog = await fetch_all(db, sql + " ORDER BY id", params)
            for ev in backlog:
                last_sent = ev["id"]
                yield {"id": str(ev["id"]), "data": json.dumps(ev)}
            while True:
                try:
                    ev = await asyncio.wait_for(queue.get(), timeout=15)
                except asyncio.TimeoutError:
                    if await request.is_disconnected():
                        break
                    continue
                if wanted(ev) and ev["id"] > last_sent:
                    last_sent = ev["id"]
                    yield {"id": str(ev["id"]), "data": json.dumps(ev)}
        finally:
            state.hub.unsubscribe(queue)

    return EventSourceResponse(generator(), ping=15)


@router.get("/business/{business_id}/stream")
async def stream_business(business_id: int, request: Request, after_id: Optional[int] = None):
    return await _stream(request, business_id, None, after_id)


@router.get("/runs/{run_id}/stream")
async def stream_run(run_id: int, request: Request, after_id: Optional[int] = None):
    async with connect() as db:
        run = await fetch_one(db, "SELECT * FROM runs WHERE id = ?", (run_id,))
    if not run:
        raise HTTPException(404, "Run not found")
    return await _stream(request, run["business_id"], run_id, after_id if after_id is not None else 0)


# ── Today tab ────────────────────────────────────────────────────────

@router.get("/business/{business_id}/overview")
async def overview(business_id: int):
    async with connect() as db:
        biz = await _business(db, business_id)
        today = await state.get_today(db)
        ingredients = await state.ingredient_rows(db, business_id, today)
        zones = await state.zone_rows(db, business_id)
        approved_docs = await state.approved_doc_types(db, business_id)
        score = calc.readiness_score([i["status"] for i in ingredients],
                                     {z["zone"]: z["status"] for z in zones}, approved_docs)
        files = await fetch_all(db, "SELECT id, filename, category, kind, zone FROM files WHERE business_id = ? "
                                    "ORDER BY id", (business_id,))
        pending = await fetch_one(db, "SELECT COUNT(*) AS n FROM approvals WHERE business_id = ? AND status = 'pending'",
                                  (business_id,))
        latest_pack = await fetch_one(db, "SELECT * FROM packs WHERE business_id = ? ORDER BY id DESC LIMIT 1",
                                      (business_id,))

    problems = [i for i in ingredients if not calc.is_covered(i["status"])]
    trays = {c: [f for f in files if f["category"] == c] for c in ("company", "products", "suppliers", "site")}
    unsorted = [f for f in files if f["category"] is None]

    journey_done = {
        "upload": bool(files),
        "check": bool(ingredients) and not unsorted,
        "fix": bool(ingredients) and not problems and all(z["status"] == "clear" for z in zones),
        "documents": len(approved_docs) >= score["documents_total"],
        "pack": latest_pack is not None,
    }
    current = next((k for k, done in journey_done.items() if not done), None)
    journey = [{"step": k, "label": k.title(), "done": v, "current": k == current} for k, v in journey_done.items()]

    renewals = [{"item": i["name"], "expiry_date": i["expiry_date"], "days": i["days_to_expiry"]}
                for i in ingredients if i["days_to_expiry"] is not None and i["days_to_expiry"] >= 0]
    if biz["halal_cert_expiry"]:
        days = calc.days_to_expiry(biz["halal_cert_expiry"], today)
        if days is not None and days >= 0:
            renewals.append({"item": "Your halal certificate", "expiry_date": biz["halal_cert_expiry"], "days": days})
    next_renewal = min(renewals, key=lambda r: r["days"]) if renewals else None

    return {
        "business": biz,
        "today": today.isoformat(),
        "readiness": score,
        "status_sentence": biz["status_sentence"] or calc.default_status_sentence(score, len(problems), pending["n"]),
        "status_by": "coordinator" if biz["status_sentence"] else "backend",
        "journey": journey,
        "trays": {c: {"count": len(fs), "files": fs[-5:]} for c, fs in trays.items()},
        "unsorted": {"count": len(unsorted), "files": unsorted[-5:]},
        "problems": [{"ingredient_id": i["id"], "name": i["name"], "status": i["status"]} for i in problems],
        "pending_approvals": pending["n"],
        "next_renewal": next_renewal,
        "latest_pack": latest_pack,
    }


# ── Agents tab ───────────────────────────────────────────────────────

@router.get("/business/{business_id}/agents")
async def agents(business_id: int):
    """Live status per agent, plus recent handoffs to draw on the graph."""
    async with connect() as db:
        last_events = {r["agent"]: r for r in await fetch_all(
            db, """SELECT e.* FROM events e JOIN (SELECT agent, MAX(id) AS id FROM events
                   WHERE business_id = ? GROUP BY agent) m ON m.id = e.id""", (business_id,))}
        counts = await fetch_all(
            db, "SELECT assignee, status, COUNT(*) AS n FROM tasks WHERE business_id = ? "
                "AND status != 'done' GROUP BY assignee, status", (business_id,))
        handoffs = await fetch_all(
            db, "SELECT id, agent, handed_to, label, task_id, created_at FROM events WHERE business_id = ? "
                "AND handed_to IS NOT NULL ORDER BY id DESC LIMIT 30", (business_id,))
    now = datetime.now(MYT)
    result = []
    for agent in AGENTS:
        open_n = sum(c["n"] for c in counts if c["assignee"] == agent and c["status"] == "open")
        busy_n = sum(c["n"] for c in counts if c["assignee"] == agent and c["status"] == "in_progress")
        last = last_events.get(agent)
        recent = last and now - datetime.fromisoformat(last["created_at"]) < timedelta(seconds=90)
        state_ = "working" if busy_n or recent else ("waiting" if open_n else "idle")
        result.append({"agent": agent, "state": state_, "open_tasks": open_n, "in_progress_tasks": busy_n,
                       "last_event": last})
    return {"agents": result, "handoffs": handoffs}


@router.get("/business/{business_id}/agents/{agent}/events")
async def agent_events(business_id: int, agent: str, limit: int = 20):
    async with connect() as db:
        rows = await fetch_all(db, "SELECT * FROM events WHERE business_id = ? AND agent = ? ORDER BY id DESC LIMIT ?",
                               (business_id, agent, min(limit, 200)))
    return list(reversed(rows))


@router.get("/business/{business_id}/feed")
async def feed(business_id: int):
    """Owner feed: Needs you / Working on / Done."""
    async with connect() as db:
        approvals = await fetch_all(
            db, """SELECT a.*, s.name AS supplier_name, s.phone AS supplier_phone, d.body AS document_body,
                          i.name AS ingredient_name
                   FROM approvals a
                   LEFT JOIN suppliers s ON s.id = a.supplier_id
                   LEFT JOIN documents d ON d.id = a.document_id
                   LEFT JOIN ingredients i ON i.id = a.ingredient_id
                   WHERE a.business_id = ? ORDER BY a.id DESC""", (business_id,))
        questions = await fetch_all(
            db, "SELECT * FROM site_findings WHERE business_id = ? AND question IS NOT NULL AND answer IS NULL "
                "AND status = 'open' ORDER BY id", (business_id,))
        tasks = await fetch_all(db, "SELECT * FROM tasks WHERE business_id = ? ORDER BY id DESC", (business_id,))

    needs_you = [{"type": "approval", **a} for a in approvals if a["status"] == "pending"]
    needs_you += [{"type": "question", **q} for q in questions]
    needs_you += [{"type": "failed_message", **a} for a in approvals if a["status"] == "failed"]
    working_on = [state.task_out(t) for t in tasks if t["status"] != "done"]
    done = [{"type": "task", "at": t["completed_at"], **state.task_out(t)} for t in tasks if t["status"] == "done"]
    done += [{"type": "message", "at": a["replied_at"] or a["sent_at"], **a}
             for a in approvals if a["status"] in ("sent", "replied")]
    done.sort(key=lambda d: d["at"] or "", reverse=True)
    return {"needs_you": needs_you, "working_on": working_on, "done": done[:30]}


@router.get("/business/{business_id}/tasks")
async def list_tasks(business_id: int, status: Optional[str] = None, assignee: Optional[str] = None):
    sql = "SELECT * FROM tasks WHERE business_id = ?"
    params: list = [business_id]
    if status:
        sql += " AND status = ?"
        params.append(status)
    if assignee:
        sql += " AND assignee = ?"
        params.append(assignee)
    async with connect() as db:
        return [state.task_out(t) for t in await fetch_all(db, sql + " ORDER BY id DESC", params)]


# ── approvals (owner actions) ────────────────────────────────────────

@router.get("/business/{business_id}/approvals")
async def list_approvals(business_id: int, status: Optional[str] = None):
    sql = "SELECT * FROM approvals WHERE business_id = ?"
    params: list = [business_id]
    if status:
        sql += " AND status = ?"
        params.append(status)
    async with connect() as db:
        return await fetch_all(db, sql + " ORDER BY id DESC", params)


class Approve(BaseModel):
    body_ms: Optional[str] = None
    body_en: Optional[str] = None


@router.post("/approvals/{approval_id}/approve")
async def approve(approval_id: int, req: Approve = Approve()):
    """One tap: approve a document, or approve and send a supplier message."""
    async with connect() as db:
        a = await fetch_one(db, "SELECT * FROM approvals WHERE id = ?", (approval_id,))
        if not a:
            raise HTTPException(404, "Approval not found")
        if a["status"] not in ("pending", "failed"):
            raise HTTPException(409, f"Already {a['status']}")
        if a["kind"] == "question":
            raise HTTPException(409, "A question is answered with POST /approvals/{id}/answer")
        now = now_iso()
        edits = {k: v for k, v in {"body_ms": req.body_ms, "body_en": req.body_en}.items() if v is not None}
        await update(db, "approvals", approval_id, {**edits, "status": "approved", "decided_at": now})
        a.update(edits)
        await db.commit()
        await state.post_event(db, a["business_id"], "owner", f"Approved: {a['title']}"[:120], "ok",
                               approval_id=approval_id, handed_to=a["requested_by"])

        if a["kind"] == "document":
            await update(db, "documents", a["document_id"], {"status": "approved", "updated_at": now})
            return await fetch_one(db, "SELECT * FROM approvals WHERE id = ?", (approval_id,))

        if a["kind"] == "plan":
            # Approved: every step of the plan becomes a task, handed over by the Coordinator.
            for step in json.loads(a["payload"] or "[]"):
                await state.create_task(db, a["business_id"], a["requested_by"], step["assignee"], step["title"][:200],
                                        detail=step.get("detail"), message=step.get("message"),
                                        run_id=a["run_id"])
            return await fetch_one(db, "SELECT * FROM approvals WHERE id = ?", (approval_id,))

        supplier = await fetch_one(db, "SELECT * FROM suppliers WHERE id = ?", (a["supplier_id"],))
        result = await messaging.deliver(a, supplier)
        if result["channel"] == "manual" and a["channel"]:
            result["channel"] = a["channel"]   # messaging is off: record it on the channel the owner chose
        await update(db, "approvals", approval_id, {**result, "sent_at": now_iso() if result["status"] == "sent" else None})
        await db.commit()
        if result["status"] == "sent":
            how = {"manual": "for you to send", "email": "by email", "whatsapp": "on WhatsApp",
                   "telegram": "on Telegram"}.get(result["channel"], f"via {result['channel']}")
            await state.post_event(db, a["business_id"], "supplier_chaser",
                                   f"Sent to {supplier['name']} {how}"[:120], "ok", approval_id=approval_id)
        else:
            await state.post_event(db, a["business_id"], "supplier_chaser",
                                   f"Could not send to {supplier['name']}"[:120], "problem", approval_id=approval_id)
        return await fetch_one(db, "SELECT * FROM approvals WHERE id = ?", (approval_id,))


class Reject(BaseModel):
    comment: Optional[str] = None


@router.post("/approvals/{approval_id}/reject")
async def reject(approval_id: int, req: Reject = Reject()):
    """Send the item back to the agent that asked, with the owner's comment as a task."""
    async with connect() as db:
        a = await fetch_one(db, "SELECT * FROM approvals WHERE id = ?", (approval_id,))
        if not a:
            raise HTTPException(404, "Approval not found")
        if a["status"] != "pending":
            raise HTTPException(409, f"Already {a['status']}")
        await update(db, "approvals", approval_id, {"status": "rejected", "comment": req.comment,
                                                    "decided_at": now_iso()})
        if a["kind"] == "document":
            await update(db, "documents", a["document_id"], {"status": "rejected", "updated_at": now_iso()})
        if a["requested_by"] in AGENTS:
            await state.create_task(
                db, a["business_id"], "owner", a["requested_by"], f"Owner asked for changes: {a['title']}"[:200],
                detail=req.comment, payload={"approval_id": approval_id, "document_id": a["document_id"],
                                             "supplier_id": a["supplier_id"]},
                event_label=f"Changes requested: {a['title']}"[:120])
        return await fetch_one(db, "SELECT * FROM approvals WHERE id = ?", (approval_id,))


@router.post("/approvals/{approval_id}/fallback")
async def force_fallback(approval_id: int):
    """Re-send a WhatsApp message on Telegram now, without waiting for the time limit."""
    async with connect() as db:
        a = await fetch_one(db, "SELECT * FROM approvals WHERE id = ?", (approval_id,))
        if not a or a["status"] != "sent":
            raise HTTPException(409, "Only a sent message can fall back")
        return await send_fallback(db, a)


async def send_fallback(db, a: dict) -> dict:
    supplier = await fetch_one(db, "SELECT * FROM suppliers WHERE id = ?", (a["supplier_id"],))
    chat_id = messaging.supplier_telegram_chat(supplier)
    now = now_iso()
    if not chat_id or not messaging.telegram_configured():
        await update(db, "approvals", a["id"], {"fallback_at": now, "send_error": "no Telegram fallback configured"})
        return {"fallback": False, "reason": "no Telegram fallback configured"}
    body = messaging.message_body(a, (supplier or {}).get("language") or "ms")
    try:
        msg_id = await messaging.send_telegram(chat_id, body)
    except Exception as exc:
        await update(db, "approvals", a["id"], {"fallback_at": now, "send_error": f"telegram: {exc}"})
        return {"fallback": False, "reason": str(exc)}
    await update(db, "approvals", a["id"], {"fallback_at": now, "channel": "telegram",
                                            "recipient": chat_id, "external_id": msg_id})
    await db.commit()
    await state.post_event(db, a["business_id"], "supplier_chaser",
                           "No WhatsApp reply; re-sent on Telegram", "attention", approval_id=a["id"])
    return {"fallback": True, "channel": "telegram"}


async def _save_uploads(business_id: int, uploads: list[UploadFile], category: str | None,
                        kind: str | None, supplier_id: int | None = None) -> list[tuple[int, str]]:
    target = UPLOAD_DIR / "files"
    target.mkdir(parents=True, exist_ok=True)
    saved = []
    async with connect() as db:
        for upload in uploads:
            name = state.safe_filename(upload.filename)
            rel = f"files/{uuid.uuid4().hex[:8]}_{name}"
            with open(UPLOAD_DIR / rel, "wb") as out:
                shutil.copyfileobj(upload.file, out)
            saved.append((await insert(db, "files", {
                "business_id": business_id, "category": category, "kind": kind, "supplier_id": supplier_id,
                "filename": name, "path": rel, "content_type": upload.content_type, "source": "upload",
                "sorted_by": "owner" if category else None, "uploaded_at": now_iso()}), name))
    return saved


@router.post("/approvals/{approval_id}/answer")
async def answer_question(approval_id: int, text: Optional[str] = Form(None),
                          files: list[UploadFile] = File(default=[])):
    """The owner answers an agent's question, with text and/or a file (e.g. a clearer copy).
    The answer goes back to the agent that asked, as a task."""
    async with connect() as db:
        a = await fetch_one(db, "SELECT * FROM approvals WHERE id = ?", (approval_id,))
    if not a or a["kind"] != "question":
        raise HTTPException(404, "Question not found")
    if a["status"] != "pending":
        raise HTTPException(409, f"Already {a['status']}")
    if not text and not files:
        raise HTTPException(422, "Give text or a file")
    supplier_id = None
    if a["ingredient_id"]:
        async with connect() as db:
            ing = await fetch_one(db, "SELECT supplier_id FROM ingredients WHERE id = ?", (a["ingredient_id"],))
        supplier_id = ing and ing["supplier_id"]
    saved = await _save_uploads(a["business_id"], files, "suppliers" if a["ingredient_id"] else None,
                                "certificate" if a["ingredient_id"] else None, supplier_id)
    async with connect() as db:
        await update(db, "approvals", approval_id, {
            "status": "replied", "reply_text": text, "reply_file_id": saved[0][0] if saved else None,
            "decided_at": now_iso(), "replied_at": now_iso()})
        ingredient = a["ingredient_id"] and await fetch_one(db, "SELECT id, name FROM ingredients WHERE id = ?",
                                                             (a["ingredient_id"],))
        await db.commit()
        await state.post_event(db, a["business_id"], "owner", f"Answered: {a['title']}"[:120], "ok",
                               approval_id=approval_id)
        task = await state.create_task(
            db, a["business_id"], "owner", a["requested_by"], f"Owner answered: {a['title']}"[:200],
            detail=text, payload={"approval_id": approval_id, "file_ids": [f for f, _ in saved],
                                  "replaces_file_id": a["file_id"],
                                  "ingredient_id": ingredient["id"] if ingredient else None,
                                  "ingredient_name": ingredient["name"] if ingredient else None},
            event_label=f"Owner sent {'a file' if saved else 'an answer'}"[:120],
            message=(text or "") + (" (file attached)" if saved else ""))
    return {"approval_id": approval_id, "task_id": task["id"], "file_ids": [f for f, _ in saved]}


# ── Records tab ──────────────────────────────────────────────────────

@router.get("/business/{business_id}/ingredients")
async def list_ingredients(business_id: int):
    async with connect() as db:
        return await state.ingredient_rows(db, business_id, await state.get_today(db))


@router.get("/ingredients/{ingredient_id}")
async def get_ingredient(ingredient_id: int):
    async with connect() as db:
        ing = await fetch_one(db, "SELECT business_id FROM ingredients WHERE id = ?", (ingredient_id,))
        if not ing:
            raise HTTPException(404, "Ingredient not found")
        today = await state.get_today(db)
        row = next(r for r in await state.ingredient_rows(db, ing["business_id"], today) if r["id"] == ingredient_id)
        certs = await fetch_all(db, "SELECT * FROM certificates WHERE ingredient_id = ? ORDER BY id DESC", (ingredient_id,))
        messages = await fetch_all(db, "SELECT * FROM approvals WHERE ingredient_id = ? ORDER BY id DESC", (ingredient_id,))
    return {**row, "certificates": certs, "messages": messages}


@router.get("/business/{business_id}/suppliers")
async def list_suppliers(business_id: int):
    async with connect() as db:
        today = await state.get_today(db)
        suppliers = await fetch_all(db, "SELECT * FROM suppliers WHERE business_id = ? ORDER BY id", (business_id,))
        ingredients = await state.ingredient_rows(db, business_id, today)
    for s in suppliers:
        s["ingredients"] = [{"id": i["id"], "name": i["name"], "status": i["status"]}
                            for i in ingredients if i["supplier_id"] == s["id"]]
    return suppliers


class PatchSupplier(BaseModel):
    contact_name: Optional[str] = None
    phone: Optional[str] = None
    email: Optional[str] = None
    telegram_chat_id: Optional[str] = None
    language: Optional[Literal["ms", "en"]] = None


@router.patch("/suppliers/{supplier_id}")
async def patch_supplier(supplier_id: int, req: PatchSupplier):
    """For the demo: point a supplier at the phone that will receive the live message."""
    async with connect() as db:
        if not await fetch_one(db, "SELECT id FROM suppliers WHERE id = ?", (supplier_id,)):
            raise HTTPException(404, "Supplier not found")
        await update(db, "suppliers", supplier_id, req.model_dump(exclude_none=True))
        return await fetch_one(db, "SELECT * FROM suppliers WHERE id = ?", (supplier_id,))


class PatchBusiness(BaseModel):
    owner_name: Optional[str] = None
    owner_phone: Optional[str] = None
    owner_telegram_chat_id: Optional[str] = None
    halal_cert_number: Optional[str] = None
    halal_cert_expiry: Optional[str] = Field(default=None, pattern=r"^\d{4}-\d{2}-\d{2}$")


@router.patch("/business/{business_id}")
async def patch_business(business_id: int, req: PatchBusiness):
    async with connect() as db:
        await _business(db, business_id)
        await update(db, "businesses", business_id, req.model_dump(exclude_none=True))
        return await _business(db, business_id)


@router.get("/business/{business_id}/documents")
async def list_documents(business_id: int):
    async with connect() as db:
        return await fetch_all(db, "SELECT * FROM documents WHERE business_id = ? ORDER BY id", (business_id,))


@router.get("/documents/{document_id}")
async def get_document(document_id: int):
    async with connect() as db:
        doc = await fetch_one(db, "SELECT * FROM documents WHERE id = ?", (document_id,))
    if not doc:
        raise HTTPException(404, "Document not found")
    return doc


class PatchDocument(BaseModel):
    body: str


@router.patch("/documents/{document_id}")
async def edit_document(document_id: int, req: PatchDocument):
    """Owner edits a draft in place (no new version; approval still needed)."""
    async with connect() as db:
        doc = await fetch_one(db, "SELECT * FROM documents WHERE id = ?", (document_id,))
        if not doc:
            raise HTTPException(404, "Document not found")
        status = "draft" if doc["status"] == "approved" else doc["status"]
        await update(db, "documents", document_id, {"body": req.body, "status": status, "updated_at": now_iso()})
        return await fetch_one(db, "SELECT * FROM documents WHERE id = ?", (document_id,))


@router.get("/business/{business_id}/site")
async def site(business_id: int):
    async with connect() as db:
        zones = await state.zone_rows(db, business_id)
    return {"zones": zones, "order": ZONES}


class PatchSiteFinding(BaseModel):
    status: Optional[Literal["open", "resolved"]] = None
    answer: Optional[str] = None


@router.patch("/site-findings/{finding_id}")
async def patch_site_finding(finding_id: int, req: PatchSiteFinding):
    async with connect() as db:
        f = await fetch_one(db, "SELECT * FROM site_findings WHERE id = ?", (finding_id,))
        if not f:
            raise HTTPException(404, "Site finding not found")
        values = req.model_dump(exclude_none=True)
        if values.get("status") == "resolved":
            values["resolved_at"] = now_iso()
        await update(db, "site_findings", finding_id, values)
        if req.answer:
            await state.create_task(
                db, f["business_id"], "owner", "onsite_inspector", f"Owner answered: {f['title']}"[:200],
                detail=req.answer, payload={"site_finding_id": finding_id, "zone": f["zone"]},
                event_label=f"Owner answered about {f['zone']}")
        return await fetch_one(db, "SELECT * FROM site_findings WHERE id = ?", (finding_id,))


# ── files ────────────────────────────────────────────────────────────

@router.post("/business/{business_id}/files")
async def upload_files(
    business_id: int,
    files: list[UploadFile] = File(...),
    category: Optional[str] = Form(None),
    kind: Optional[str] = Form(None),
    zone: Optional[str] = Form(None),
):
    """Drop files into a tray (category), or with no category for the Coordinator to sort."""
    if category and category not in ("company", "products", "suppliers", "site"):
        raise HTTPException(422, "Unknown category")
    if zone and zone not in ZONES:
        raise HTTPException(422, "Unknown zone")
    target = UPLOAD_DIR / "files"
    target.mkdir(parents=True, exist_ok=True)
    saved = []
    async with connect() as db:
        await _business(db, business_id)
        for upload in files:
            name = state.safe_filename(upload.filename)
            rel = f"files/{uuid.uuid4().hex[:8]}_{name}"
            with open(UPLOAD_DIR / rel, "wb") as out:
                shutil.copyfileobj(upload.file, out)
            row = {"business_id": business_id, "category": category or None, "kind": kind or None,
                   "zone": zone or None, "filename": name, "path": rel, "content_type": upload.content_type,
                   "source": "upload", "uploaded_at": now_iso()}
            row["id"] = await insert(db, "files", row)
            saved.append(row)
        await db.commit()
        label = f"Added {saved[0]['filename']}" if len(saved) == 1 else f"Added {len(saved)} files"
        await state.post_event(db, business_id, "owner", label[:120], "info",
                               handed_to=None if category else "coordinator")
    return {"files": saved}


@router.get("/business/{business_id}/files")
async def list_files(business_id: int, category: Optional[str] = None):
    async with connect() as db:
        if category == "unsorted":
            return await fetch_all(db, "SELECT * FROM files WHERE business_id = ? AND category IS NULL", (business_id,))
        if category:
            return await fetch_all(db, "SELECT * FROM files WHERE business_id = ? AND category = ?",
                                   (business_id, category))
        return await fetch_all(db, "SELECT * FROM files WHERE business_id = ? ORDER BY id", (business_id,))


def _thumbnail(path, filename: str):
    """A small JPEG of a file: photos shrunk, a scanned PDF's page image, a drawn grid of a sheet's first rows."""
    import io
    import re as _re
    from PIL import Image, ImageDraw, ImageFont
    name = filename.lower()
    img = None
    if name.endswith((".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp")):
        img = Image.open(path)
    elif name.endswith(".pdf"):
        data = path.read_bytes()
        m = _re.search(rb"/DCTDecode.*?stream\r?\n", data, _re.S)
        if m:
            end = data.find(b"endstream", m.end())
            img = Image.open(io.BytesIO(data[m.end():end]))
    elif name.endswith(".xlsx"):
        from openpyxl import load_workbook
        rows = [["" if v is None else str(v) for v in r][:4]
                for r in load_workbook(path, read_only=True, data_only=True).worksheets[0].iter_rows(max_row=9, values_only=True)]
        img = Image.new("RGB", (360, 470), "white")
        d = ImageDraw.Draw(img)
        try:
            font = ImageFont.truetype("arial.ttf", 13)
            bold = ImageFont.truetype("arialbd.ttf", 13)
        except OSError:
            font = bold = ImageFont.load_default()
        d.rectangle([0, 0, 360, 34], fill=(33, 115, 70))
        d.text((12, 10), filename[:40], font=bold, fill="white")
        for i, row in enumerate(rows):
            y = 46 + i * 44
            if i == 0:
                d.rectangle([0, y - 6, 360, y + 30], fill=(236, 244, 239))
            for k, cell in enumerate(row):
                d.text((10 + k * 88, y + 4), cell[:11], font=bold if i == 0 else font, fill=(30, 30, 34))
            d.line([0, y + 30, 360, y + 30], fill=(220, 220, 226))
        for k in range(1, 4):
            d.line([k * 88 + 2, 40, k * 88 + 2, 470], fill=(230, 230, 234))
    if img is None:
        return None
    img = img.convert("RGB")
    img.thumbnail((360, 480))
    out = io.BytesIO()
    img.save(out, "JPEG", quality=78)
    return out.getvalue()


@router.get("/files/{file_id}/thumb")
async def file_thumb(file_id: int):
    """A small picture of the file for the Input step; 404 when there is nothing to show (the page draws an icon)."""
    from fastapi.responses import Response
    async with connect() as db:
        f = await fetch_one(db, "SELECT * FROM files WHERE id = ?", (file_id,))
    if not f or not pack.file_path(f).exists():
        raise HTTPException(404, "File not found")
    cache = UPLOAD_DIR / "thumbs" / f"{file_id}.jpg"
    if not cache.exists():
        try:
            data = _thumbnail(pack.file_path(f), f["filename"])
        except Exception:
            data = None
        if not data:
            raise HTTPException(404, "No picture for this file")
        cache.parent.mkdir(parents=True, exist_ok=True)
        cache.write_bytes(data)
    return Response(cache.read_bytes(), media_type="image/jpeg", headers={"Cache-Control": "max-age=3600"})


@router.get("/files/{file_id}/sheet")
async def file_sheet(file_id: int, limit: int = 60):
    """The first sheet of an .xlsx file as rows, for the file viewer."""
    from openpyxl import load_workbook
    async with connect() as db:
        f = await fetch_one(db, "SELECT * FROM files WHERE id = ?", (file_id,))
    if not f or not pack.file_path(f).exists() or not f["filename"].lower().endswith(".xlsx"):
        raise HTTPException(404, "Spreadsheet not found")
    wb = load_workbook(pack.file_path(f), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = [["" if v is None else str(v) for v in r] for r in ws.iter_rows(max_row=limit, values_only=True)]
    while rows and not any(rows[-1]):
        rows.pop()
    return {"sheet": ws.title, "sheets": wb.sheetnames, "rows": rows}


class AskSupplier(BaseModel):
    ingredient_id: int
    channel: Literal["whatsapp", "telegram", "email"]


@router.post("/business/{business_id}/ask-supplier")
async def ask_supplier(business_id: int, req: AskSupplier):
    """The owner asks the Supplier Chaser to write to an ingredient's supplier on the channel she picked."""
    async with connect() as db:
        ing = await fetch_one(db, """SELECT i.*, s.name AS supplier_name, s.contact_name FROM ingredients i
                                     LEFT JOIN suppliers s ON s.id = i.supplier_id
                                     WHERE i.id = ? AND i.business_id = ?""", (req.ingredient_id, business_id))
        if not ing or not ing["supplier_name"]:
            raise HTTPException(404, "No supplier on file for this ingredient")
        how = {"whatsapp": "a WhatsApp", "telegram": "a Telegram message", "email": "an email"}[req.channel]
        task = await state.create_task(
            db, business_id, "owner", "supplier_chaser",
            f"Ask {ing['supplier_name']} for the renewed {ing['name'].lower()} certificate",
            payload={"channel": req.channel, "supplier_name": ing["supplier_name"], "ingredient_name": ing["name"]},
            message=f"Please send {how} to {ing['contact_name'] or ing['supplier_name']} for the renewed certificate.")
    return {"task_id": task["id"]}


@router.post("/business/{business_id}/follow-up")
async def follow_up(business_id: int, key: str = Form(...), files: list[UploadFile] = File(...)):
    """The owner uploads a fix for one open checklist row (a renewed certificate, a new zone photo).
    It goes to the agent that owns that row, as a task; any question about the same thing is closed."""
    kind, _, ref = key.partition(":")
    async with connect() as db:
        await _business(db, business_id)
        if kind == "cert":
            ing = await fetch_one(db, "SELECT * FROM ingredients WHERE id = ? AND business_id = ?", (int(ref), business_id))
            if not ing:
                raise HTTPException(404, "Ingredient not found")
            saved = await _save_uploads(business_id, files, "suppliers", "certificate", ing["supplier_id"])
            assignee, title = "certificate_checker", f"Owner uploaded a certificate for {ing['name']}"
            payload = {"ingredient_id": ing["id"], "ingredient_name": ing["name"]}
            open_q = "kind = 'question' AND ingredient_id = ?", (ing["id"],)
        elif kind == "zone" and ref in ZONES:
            saved = await _save_uploads(business_id, files, "site", "photo")
            for fid, _ in saved:
                await update(db, "files", fid, {"zone": ref})
            assignee, title = "onsite_inspector", f"Owner uploaded a new photo of the {ref} zone"
            payload = {"zone": ref}
            open_q = "kind = 'question' AND payload LIKE ?", (f'%"zone": "{ref}"%',)
        else:
            raise HTTPException(422, "This row has no upload")
        payload["file_ids"] = [fid for fid, _ in saved]
        now = now_iso()
        await db.execute(f"UPDATE approvals SET status = 'replied', reply_file_id = ?, decided_at = ?, replied_at = ? "
                         f"WHERE business_id = ? AND status = 'pending' AND {open_q[0]}",
                         (saved[0][0], now, now, business_id, *open_q[1]))
        await db.commit()
        task = await state.create_task(db, business_id, "owner", assignee, title, payload=payload,
                                       event_label=f"Uploaded {saved[0][1]}"[:120],
                                       message=f"I uploaded {saved[0][1]}. Can you check it?")
    return {"task_id": task["id"], "file_ids": payload["file_ids"]}


@router.get("/files/{file_id}/content")
async def file_content(file_id: int):
    async with connect() as db:
        f = await fetch_one(db, "SELECT * FROM files WHERE id = ?", (file_id,))
    if not f or not pack.file_path(f).exists():
        raise HTTPException(404, "File not found")
    if f.get("sha256") and fingerprint(pack.file_path(f)) != f["sha256"]:
        raise HTTPException(409, f"{f['filename']} changed after it was uploaded; it can't be trusted")
    return FileResponse(pack.file_path(f), media_type=f["content_type"], filename=f["filename"],
                        content_disposition_type="inline")


# ── pack ─────────────────────────────────────────────────────────────

@router.get("/business/{business_id}/pack")
async def pack_checklist(business_id: int):
    async with connect() as db:
        await _business(db, business_id)
        return await pack.checklist(db, business_id)


@router.post("/business/{business_id}/pack")
async def pack_assemble(business_id: int):
    async with connect() as db:
        await _business(db, business_id)
        return await pack.assemble(db, business_id)


@router.get("/business/{business_id}/pack/download")
async def pack_download(business_id: int):
    async with connect() as db:
        latest = await fetch_one(db, "SELECT * FROM packs WHERE business_id = ? ORDER BY id DESC LIMIT 1",
                                 (business_id,))
    if not latest or not (UPLOAD_DIR / latest["path"]).exists():
        raise HTTPException(404, "No pack assembled yet")
    return FileResponse(UPLOAD_DIR / latest["path"], media_type="application/zip",
                        filename=latest["path"].split("/")[-1])


# ── upkeep and messaging status ──────────────────────────────────────

@router.post("/business/{business_id}/upkeep/check")
async def upkeep_check(business_id: int):
    async with connect() as db:
        await _business(db, business_id)
        created = await state.upkeep_check(db, business_id)
    return {"triggered_tasks": [state.task_out(t) for t in created]}


@router.get("/messaging/status")
async def messaging_status():
    async with connect() as db:
        last_chat = await fetch_one(db, "SELECT value FROM settings WHERE key = 'telegram_last_chat'")
        url = await fetch_one(db, "SELECT value FROM settings WHERE key = 'public_webhook_url'")
    result = {**messaging.status(), "telegram_last_chat": json.loads(last_chat["value"]) if last_chat else None}
    if url and url["value"]:
        result["public_webhook_url"] = url["value"]
    return result


class PublicUrl(BaseModel):
    url: str = Field(pattern=r"^https://")


@router.put("/messaging/public-url")
async def set_public_url(req: PublicUrl):
    """Called by scripts/tunnel.py (locally) with the ngrok address, used for Twilio signature checks."""
    async with connect() as db:
        await db.execute("INSERT INTO settings (key, value) VALUES ('public_webhook_url', ?) "
                         "ON CONFLICT(key) DO UPDATE SET value = excluded.value", (req.url.rstrip("/"),))
    return {"public_webhook_url": req.url.rstrip("/"),
            "twilio_webhook": req.url.rstrip("/") + "/webhooks/twilio",
            "telegram_webhook": req.url.rstrip("/") + "/webhooks/telegram"}
