r"""
Generate the files for the scripted demo run (Mak Teh Sambal, fictional).

    .venv\Scripts\python scripts/make_demo_case.py

Writes demo_case/:
  upload/      the owner's files, dropped in all at once on the Input step
  follow_up/   what the owner uploads later from the output checklist
               (the renewed oil certificate, the cleaning area after the fix)
  case.json    what the scripted agents "read" in each file, with the position
               of every field so the dashboard can point at it

Drawing helpers come from make_sample_case.py. Every file carries a small
"SAMPLE - FICTIONAL" footer instead of the large red band.
"""

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageDraw

sys.path.insert(0, str(Path(__file__).resolve().parent))
import make_sample_case as msc  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "demo_case"
UPLOAD = OUT / "upload"
FOLLOW = OUT / "follow_up"

FOOTER = "SAMPLE - FICTIONAL"
TODAY = "2026-10-15"

# The big red band and the diagonal watermark give way to a small footer.
msc.sample_band = lambda d, w, text=None: None
msc.watermark = lambda img, text=None: None


def footer(img: Image.Image):
    d = ImageDraw.Draw(img)
    d.text((img.width - 24, img.height - 22), FOOTER, font=msc.font(15), fill=(150, 150, 156), anchor="rs")


def photo_caption(img: Image.Image, text: str):
    d = ImageDraw.Draw(img)
    d.rectangle([0, 960, 1600, 1000], fill=(0, 0, 0))
    d.text((20, 980), text, font=msc.font(20), fill="white", anchor="lm")
    d.text((1580, 980), FOOTER, font=msc.font(15), fill=(170, 170, 170), anchor="rm")


msc.photo_caption = photo_caption

COMPANY = msc.COMPANY
JAKIM, JAIS = msc.JAKIM, msc.JAIS

SUPPLIERS = [
    # name, contact, phone, supplies, brand
    ("Syarikat Cili Maju", "Encik Halim", "+60 11-0000 0101", "Dried chilli", "CILI MAJU"),
    ("Minyak Murni Sdn Bhd", "Puan Aishah", "+60 11-0000 0202", "Palm cooking oil, 17 kg tin", "MURNI"),
    ("Bawang Segar Enterprise", "Encik Kumar", "+60 11-0000 0303", "Shallot paste, garlic paste", "BAWANG SEGAR"),
    ("Hasil Laut Kuala", "Puan Zainab", "+60 11-0000 0404", "Ikan bilis (dried anchovies)", "HASIL LAUT"),
    ("Gula Manis Sdn Bhd", "Encik Lim", "+60 11-0000 0505", "Fine sugar, 1 kg bag", "GULA MANIS"),
]

INGREDIENTS = [
    # name, brand, supplier, kind, used in
    ("Dried chilli", "CILI MAJU", "Syarikat Cili Maju", "processed", ["Sambal Tumis", "Sambal Ikan Bilis"]),
    ("Cooking oil", "MURNI", "Minyak Murni Sdn Bhd", "processed", ["Sambal Tumis", "Sambal Ikan Bilis"]),
    ("Shallot paste", "BAWANG SEGAR", "Bawang Segar Enterprise", "processed", ["Sambal Tumis", "Sambal Ikan Bilis"]),
    ("Garlic paste", "BAWANG SEGAR", "Bawang Segar Enterprise", "processed", ["Sambal Tumis", "Sambal Ikan Bilis"]),
    ("Ikan bilis", "HASIL LAUT", "Hasil Laut Kuala", "processed", ["Sambal Ikan Bilis"]),
    ("Sugar", "GULA MANIS", "Gula Manis Sdn Bhd", "processed", ["Sambal Tumis", "Sambal Ikan Bilis"]),
]

CERTS = {
    "chilli": dict(holder="Syarikat Cili Maju", products=["Cili Kering (Dried Chilli) - jenama CILI MAJU"],
                   number="HC-CM-2025-0412", issued="1 Sep 2025", expiry="31 Aug 2027", expiry_iso="2027-08-31",
                   issuer=JAKIM, ingredients=["Dried chilli"]),
    "shallot_garlic": dict(holder="Bawang Segar Enterprise",
                           products=["Pes Bawang Merah (Shallot Paste) - jenama BAWANG SEGAR",
                                     "Pes Bawang Putih (Garlic Paste) - jenama BAWANG SEGAR"],
                           number="HC-BS-2025-0733", issued="1 Mar 2025", expiry="28 Feb 2027",
                           expiry_iso="2027-02-28", issuer=JAKIM, ingredients=["Shallot paste", "Garlic paste"]),
    "oil": dict(holder="Minyak Murni Sdn Bhd", products=["Minyak Masak Sawit (Palm Cooking Oil) - jenama MURNI"],
                number="HC-MM-2024-0587", issued="1 Sep 2024", expiry="31 Aug 2026", expiry_iso="2026-08-31",
                issuer=JAKIM, ingredients=["Cooking oil"]),
    "ikan_bilis": dict(holder="Hasil Laut Kuala", products=["Ikan Bilis Kering (Dried Anchovies) - jenama HASIL LAUT"],
                       number="JAIS-HLK-2025-0158", issued="1 Oct 2025", expiry="30 Sep 2027",
                       expiry_iso="2027-09-30", issuer=JAIS, ingredients=["Ikan bilis"]),
    "sugar": dict(holder="Gula Manis Sdn Bhd", products=["Gula Pasir Halus (Fine Sugar) - jenama GULA MANIS"],
                  number="HC-GM-2025-0961", issued="1 Apr 2025", expiry="31 Mar 2027", expiry_iso="2027-03-31",
                  issuer=JAKIM, ingredients=["Sugar"]),
    # follow-up: the renewed oil certificate
    "oil_renewed": dict(holder="Minyak Murni Sdn Bhd", products=["Minyak Masak Sawit (Palm Cooking Oil) - jenama MURNI"],
                        number="HC-MM-2026-0912", issued="1 Sep 2026", expiry="31 Aug 2028", expiry_iso="2028-08-31",
                        issuer=JAKIM, ingredients=["Cooking oil"]),
}


# ── certificates, with the box of every field ────────────────────────

def certificate(key: str) -> tuple[Image.Image, dict]:
    c = CERTS[key]
    img, d = msc.page()
    w, h = img.size
    green = (60, 110, 70)
    d.rectangle([40, 40, w - 40, h - 40], outline=green, width=6)
    d.rectangle([54, 54, w - 54, h - 54], outline=green, width=2)
    msc.leaf_mark(d, w // 2, 160, 60, green)
    d.text((w // 2, 260), c["issuer"].upper(), font=msc.font(26, True), fill=msc.INK, anchor="mm")
    d.text((w // 2, 330), "SIJIL PENGESAHAN HALAL", font=msc.font(46, True), fill=(40, 90, 50), anchor="mm")
    d.text((w // 2, 380), "HALAL CERTIFICATE", font=msc.font(24), fill=msc.GREY, anchor="mm")
    boxes = {"title": [w // 2 - 330, 300, w // 2 + 330, 400]}

    x0, cols, size = 100, [340, 560], 24
    f, fb = msc.font(size), msc.font(size, True)
    rows = [("Perkara / Field", "Butiran / Details", None),
            ("Nama Syarikat / Company", c["holder"], "holder")]
    rows += [("Produk / Product(s)" if i == 0 else "", p, "products") for i, p in enumerate(c["products"])]
    rows += [("No. Sijil / Certificate No.", c["number"], "number"),
             ("Tarikh Dikeluarkan / Issued", c["issued"], "issued"),
             ("Sah Hingga / Valid Until", c["expiry"], "expiry"),
             ("Dikeluarkan oleh / Issued by", c["issuer"], "issuer")]
    y = 450
    for r, (label, value, key_) in enumerate(rows):
        cells = [msc.wrap(d, label, fb if r == 0 else f, cols[0] - 16), msc.wrap(d, value, fb if r == 0 else f, cols[1] - 16)]
        rh = max(len(x) for x in cells) * (size + 8) + 16
        cx = x0
        for i, lines in enumerate(cells):
            if r == 0:
                d.rectangle([cx, y, cx + cols[i], y + rh], fill=(236, 234, 228))
            d.rectangle([cx, y, cx + cols[i], y + rh], outline=msc.LINE, width=2)
            for k, line in enumerate(lines):
                d.text((cx + 8, y + 8 + k * (size + 8)), line, font=fb if r == 0 else f, fill=msc.INK)
            cx += cols[i]
        if key_:
            box = [x0 + cols[0], y, x0 + cols[0] + cols[1], y + rh]
            old = boxes.get(key_)
            boxes[key_] = [old[0], old[1], box[2], box[3]] if old else box
        y += rh
    d.text((100, y + 50), "Sijil ini sah bagi produk yang disenaraikan sahaja.", font=msc.font(22), fill=msc.INK)
    d.text((100, y + 82), "This certificate is valid only for the products listed above.", font=msc.font(22), fill=msc.GREY)
    d.line([680, y + 220, 1000, y + 220], fill=msc.INK, width=2)
    d.text((840, y + 240), "Pengarah / Director", font=msc.font(20), fill=msc.GREY, anchor="mm")
    footer(img)
    frac = {k: [round(b[0] / w, 4), round(b[1] / h, 4), round(b[2] / w, 4), round(b[3] / h, 4)] for k, b in boxes.items()}
    return img, frac


def scan(img: Image.Image, seed: int, angle: float) -> Image.Image:
    """A phone photo of a paper certificate: a slight tilt and some grain. Still easy to read."""
    out = img.rotate(angle, resample=Image.BICUBIC, fillcolor=(214, 210, 200), expand=False)
    return msc.photo_texture(out, seed).convert("RGB")


# ── spreadsheets ─────────────────────────────────────────────────────

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="2F5D3A")


def sheet(path: Path, title: str, header: list[str], rows: list[list], widths: list[int]):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HEAD, HEAD_FILL, Alignment(vertical="center")
    for r in rows:
        ws.append(r)
    for i, wdt in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = wdt
    ws.freeze_panes = "A2"
    ws.cell(row=len(rows) + 3, column=1, value=FOOTER).font = Font(size=8, color="999999")
    wb.save(path)


# ── photos ───────────────────────────────────────────────────────────

def store_photo(fixed: bool) -> Image.Image:
    """The store shelf. fixed=True: the two Cap Merpati Hijau tins are gone, two MURNI tins in their place."""
    real_tin = msc.tin

    def tin(d, x, y, brand, line2, color, text_color="white", w=150, h=190):
        if fixed and brand.startswith("CAP"):
            return real_tin(d, x, y, "MURNI", "Minyak Masak Sawit", (230, 180, 30), text_color=(90, 50, 0), w=w, h=h)
        return real_tin(d, x, y, brand, line2, color, text_color, w, h)

    msc.tin = tin
    try:
        img = msc.store_photo()
    finally:
        msc.tin = real_tin
    if fixed:
        photo_caption(img, "Site photo: raw material store, 15 Oct 2026, 11:40")
    return img


def cleaning_photo(detergent: bool) -> Image.Image:
    """The cleaning area. detergent=True: a detergent bottle stands on the counter next to a tub of chilli paste."""
    img = msc.cleaning_photo()
    if detergent:
        d = ImageDraw.Draw(img)
        d.rounded_rectangle([200, 395, 330, 470], 8, fill=(250, 250, 248), outline=(90, 90, 90), width=3)
        d.rectangle([194, 382, 336, 398], fill=(200, 60, 40))
        d.text((265, 422), "PES CILI", font=msc.font(20, True), fill=(160, 30, 20), anchor="mm")
        d.text((265, 448), "untuk masakan", font=msc.font(13), fill=msc.INK, anchor="mm")
        d.rounded_rectangle([345, 345, 395, 470], 10, fill=(40, 120, 200), outline=(20, 60, 110), width=3)
        d.rectangle([358, 325, 382, 347], fill=(230, 230, 230), outline=(90, 90, 90), width=2)
        d.text((370, 390), "SABUN", font=msc.font(12, True), fill="white", anchor="mm")
        d.text((370, 408), "PINGGAN", font=msc.font(10, True), fill="white", anchor="mm")
        d.text((370, 432), "Detergent", font=msc.font(10), fill="white", anchor="mm")
    photo_caption(img, "Site photo: cleaning area, " + ("10 Oct 2026, 09:33" if detergent else "15 Oct 2026, 11:52"))
    return img


ZONES = {
    "store": {"title": "Raw material store", "checks": [
        {"label": "Items labelled", "state": "ok", "detail": "Every sack, tub and tin is labelled."},
        {"label": "Stored off the floor", "state": "ok", "detail": "Everything is on the rack or the pallet."}]},
    "processing": {"title": "Processing area", "checks": [
        {"label": "Clean surfaces", "state": "ok", "detail": "Clean steel tables, hygiene sign on the wall."}]},
    "packing": {"title": "Packing area", "checks": [
        {"label": "Jars sealed and coded", "state": "ok", "detail": "All 10 jars have lids; batch code check sign above the table."}]},
    "cleaning": {"title": "Cleaning area", "checks": [
        {"label": "Detergent next to food", "state": "warn", "ask": "Please move it",
         "detail": "A bottle of dishwashing detergent stands next to the tub of chilli paste on the counter.",
         "box": [0.115, 0.315, 0.255, 0.48]},
        {"label": "Chemicals cabinet closed", "state": "ok", "detail": "Red cabinet marked 'Cleaning chemicals only'."}],
        "suggestion": "Please move the detergent away from the food, into the chemicals cabinet. Then send a new photo "
                      "of the counter.",
        "after": [{"label": "Detergent stored away from food", "state": "ok", "detail": "Nothing but food items on the counter."},
                  {"label": "Chemicals cabinet closed", "state": "ok", "detail": "Red cabinet marked 'Cleaning chemicals only'."}]},
    "waste": {"title": "Waste area", "checks": [
        {"label": "Bins closed", "state": "ok", "detail": "All 3 bins have their lids on, outside the back door."}]},
}


# ── documents ────────────────────────────────────────────────────────

DOCUMENTS = {
    "halal_policy": ("Halal Policy", f"""## Policy statement
{COMPANY['name']} produces Sambal Tumis Mak Teh and Sambal Ikan Bilis Mak Teh as halal food. We use only ingredients with a valid halal certificate from a body recognised by JAKIM, and we keep our premises free of anything that is not halal.

## Scope
This policy covers both products, every ingredient and supplier, the factory unit at {COMPANY['address']}, and all six production and packing staff.

## Halal supervisor
{COMPANY['supervisor']} is the halal supervisor. She checks incoming ingredients, keeps the certificate register up to date and briefs new staff before they start work.

## Commitments
- Buy only from the approved supplier list, and only brands named on it
- Keep every supplier's halal certificate on file and renew it before it expires
- Keep anything not on the ingredient list out of the store and the production area
- Train staff on halal handling when they join and every six months
- Review this policy once a year, or whenever a product, supplier or process changes

## Approval
Approved by {COMPANY['owner']}, Managing Director. To be signed and dated before submission.
"""),
    "ingredient_control": ("Ingredient Control Procedure", f"""## Purpose
Make sure every ingredient used in Mak Teh sambal comes from an approved supplier and has a valid halal certificate on file.

## Approved suppliers
Ingredients are bought only from the five suppliers on the approved supplier list (updated 1 Oct 2026): Syarikat Cili Maju, Minyak Murni Sdn Bhd, Bawang Segar Enterprise, Hasil Laut Kuala and Gula Manis Sdn Bhd. A new supplier is added only after its certificate has been checked.

## Receiving checks
At every delivery the store keeper checks that the brand matches the approved list, that the halal logo and certificate number are on the pack, and that the packaging is intact. The supplier's lot number is written in the receiving log.

## Certificate monitoring
The halal supervisor keeps a register of every certificate with its number, issuer and expiry date. Sixty days before a certificate expires, the supplier is asked for the renewed one.

## Non-conforming materials
Anything without a valid certificate, or a brand that is not on the list, is kept apart, labelled 'Do not use' and reported to the owner the same day. It is returned or thrown away, never used.

## Records
The receiving log, the certificate register and copies of all certificates are kept for at least three years.
"""),
    "traceability": ("Traceability Procedure", """## Purpose
Make it possible to trace any jar back to the ingredient lots it was made from, and every ingredient lot forward to the jars that used it.

## Batch coding
Every batch gets a code: product letter, date and batch number (for example T-261015-01). The code is printed on every jar label.

## Records
For each batch the production lead records the date, the batch code, the lot number of every ingredient used, the staff on the shift and the number of jars filled.

## Tracing a product
Starting from the batch code on a jar, the batch record shows the ingredient lots; the receiving log shows which delivery and supplier each lot came from.

## Recall
If an ingredient is found not to be halal, every batch that used its lot is identified from the records, held, and recalled from customers. The owner informs JAKIM within 24 hours.
"""),
}


# ── write everything ─────────────────────────────────────────────────

def main():
    for d in (UPLOAD, FOLLOW):
        d.mkdir(parents=True, exist_ok=True)
        for old in d.glob("*"):
            old.unlink()

    files, certs = [], []

    def save_pdf(img: Image.Image, name: str, category: str, kind: str):
        footer(img)
        img.convert("RGB").save(UPLOAD / name, "PDF", resolution=150)
        files.append({"filename": name, "category": category, "kind": kind})

    save_pdf(msc.ssm_registration(), "01_SSM_registration.pdf", "company", "registration")
    save_pdf(msc.company_profile(), "02_Company_profile.pdf", "company", "other")

    sheet(UPLOAD / "03_Ingredient_list.xlsx", "Ingredients",
          ["Ingredient", "Brand", "Supplier", "Halal proof", "Used in"],
          [[n, b, s, "Halal certificate", ", ".join(p)] for n, b, s, _, p in INGREDIENTS], [18, 16, 26, 18, 34])
    files.append({"filename": "03_Ingredient_list.xlsx", "category": "products", "kind": "other"})
    sheet(UPLOAD / "04_Supplier_list.xlsx", "Approved suppliers",
          ["Supplier", "Contact", "Phone (WhatsApp)", "Supplies", "Brand"], [list(s) for s in SUPPLIERS],
          [26, 14, 18, 30, 16])
    files.append({"filename": "04_Supplier_list.xlsx", "category": "suppliers", "kind": "other"})

    tumis = [("Cili kering (dried chilli)", "CILI MAJU", "3 kg"),
             ("Minyak masak sawit (palm cooking oil)", "MURNI", "4 kg"),
             ("Pes bawang merah (shallot paste)", "BAWANG SEGAR", "2 kg"),
             ("Pes bawang putih (garlic paste)", "BAWANG SEGAR", "0.5 kg"),
             ("Gula (sugar)", "GULA MANIS", "1.2 kg")]
    bilis = tumis + [("Ikan bilis (dried anchovies)", "HASIL LAUT", "2 kg")]
    save_pdf(msc.recipe("Sambal Tumis Mak Teh (250 g)", tumis, [
        "Soak and blend the dried chilli.", "Heat the oil; fry shallot and garlic paste until fragrant.",
        "Add the chilli and cook on low heat for 45 minutes, stirring.",
        "Add sugar, cook 10 more minutes, then fill hot into jars."]), "05_Recipe_Sambal_Tumis.pdf", "products", "recipe")
    save_pdf(msc.recipe("Sambal Ikan Bilis Mak Teh (250 g)", bilis, [
        "Fry the ikan bilis in part of the oil until crisp; set aside.", "Cook the sambal base as for Sambal Tumis.",
        "Fold the ikan bilis into the sambal and fill hot into jars."]), "06_Recipe_Sambal_Ikan_Bilis.pdf", "products", "recipe")

    for n, (key, name, supplier, angle) in enumerate([
            ("chilli", "07_Halal_cert_Cili_Maju.jpg", "Syarikat Cili Maju", -1.2),
            ("shallot_garlic", "08_Halal_cert_Bawang_Segar.jpg", "Bawang Segar Enterprise", 0.8),
            ("oil", "09_Halal_cert_Minyak_Murni.jpg", "Minyak Murni Sdn Bhd", -0.6),
            ("ikan_bilis", "10_Halal_cert_Hasil_Laut.jpg", "Hasil Laut Kuala", 1.1),
            ("sugar", "11_Halal_cert_Gula_Manis.jpg", "Gula Manis Sdn Bhd", -0.9)]):
        img, boxes = certificate(key)
        scan(img, 70 + n, angle).save(UPLOAD / name, "JPEG", quality=86)
        files.append({"filename": name, "category": "suppliers", "kind": "certificate", "supplier": supplier})
        certs.append({"file": name, "key": key, "supplier": supplier, **CERTS[key], "boxes": boxes})

    photos = [("12_Photo_store.jpg", "store", store_photo(True)),
              ("13_Photo_processing.jpg", "processing", msc.processing_photo()),
              ("14_Photo_packing.jpg", "packing", msc.packing_photo()),
              ("15_Photo_cleaning.jpg", "cleaning", cleaning_photo(True)),
              ("16_Photo_waste.jpg", "waste", msc.waste_photo())]
    for k, (name, zone, img) in enumerate(photos):
        msc.photo_texture(img, 12 + k).convert("RGB").save(UPLOAD / name, "JPEG", quality=84)
        files.append({"filename": name, "category": "site", "kind": "photo", "zone": zone})

    # follow-ups the owner uploads from the output checklist
    img, boxes = certificate("oil_renewed")
    scan(img, 99, 0.7).save(FOLLOW / "Halal_cert_Minyak_Murni_2026.jpg", "JPEG", quality=86)
    renewed = {"file": "Halal_cert_Minyak_Murni_2026.jpg", "key": "oil_renewed", "supplier": "Minyak Murni Sdn Bhd",
               **CERTS["oil_renewed"], "boxes": boxes}
    msc.photo_texture(cleaning_photo(False), 21).convert("RGB").save(FOLLOW / "Photo_cleaning_after.jpg", "JPEG", quality=84)

    case = {"today": TODAY, "company": COMPANY, "suppliers": SUPPLIERS, "ingredients": INGREDIENTS,
            "files": files, "certificates": certs, "renewed_certificate": renewed, "zones": ZONES,
            "documents": {k: {"title": t, "body": b} for k, (t, b) in DOCUMENTS.items()}}
    (OUT / "case.json").write_text(json.dumps(case, indent=1, ensure_ascii=False), encoding="utf-8")
    print(f"Wrote {len(files)} files to {UPLOAD.relative_to(ROOT)}, 2 to {FOLLOW.relative_to(ROOT)} and case.json")


if __name__ == "__main__":
    main()
